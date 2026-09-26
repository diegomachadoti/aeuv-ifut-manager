from __future__ import annotations

import argparse
import configparser
import logging
import re
import sys
import time
import unicodedata
import io
from datetime import datetime
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable
from urllib.parse import urlparse

import requests
from google.oauth2 import service_account
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload, MediaFileUpload
from selenium import webdriver
from selenium.common.exceptions import NoSuchElementException, TimeoutException
from selenium.webdriver import ChromeOptions
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.common.by import By
from selenium.webdriver.remote.webelement import WebElement
from selenium.webdriver.support import expected_conditions as ec
from selenium.webdriver.support.ui import WebDriverWait
from webdriver_manager.chrome import ChromeDriverManager


DEFAULT_CONFIG_PATH = Path("config.ini")
DEFAULT_DOWNLOAD_DIR = Path("downloads")
DEFAULT_LOG_PATH = Path("logs") / "ifut.log"
DEFAULT_PROCESSED_DIR = Path("downloads") / "processados"
DEFAULT_FAILED_DIR = Path("downloads") / "falhas"
DEFAULT_SELECTORS_PATH = Path("selectors.ini")


def normalize_text(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", value)
    normalized = normalized.encode("ascii", "ignore").decode("ascii")
    normalized = normalized.lower()
    normalized = re.sub(r"[^a-z0-9]+", " ", normalized)
    return normalized.strip()


def team_search_key(team_name: str) -> str:
    return normalize_text(team_name.split("-", 1)[0])


def text_tokens(text: str) -> set[str]:
    """Extrai tokens significantes ignorando particulas comuns."""
    particles = {"da", "de", "do", "a", "o", "e", "em"}
    normalized = normalize_text(text)
    tokens = {token for token in normalized.split() if token and token not in particles}
    return tokens


def text_tokens_all(text: str) -> set[str]:
    normalized = normalize_text(text)
    return {token for token in normalized.split() if token}


def token_overlap_score(target_tokens: set[str], candidate_tokens: set[str]) -> float:
    """Calcula similaridade entre conjuntos de tokens (0.0 a 1.0)."""
    if not target_tokens or not candidate_tokens:
        return 0.0
    intersection = len(target_tokens & candidate_tokens)
    union = len(target_tokens | candidate_tokens)
    return intersection / union if union > 0 else 0.0


def token_containment_score(target_tokens: set[str], candidate_tokens: set[str]) -> float:
    """Calcula quanto o conjunto menor esta contido no maior."""
    if not target_tokens or not candidate_tokens:
        return 0.0
    smaller = target_tokens if len(target_tokens) <= len(candidate_tokens) else candidate_tokens
    larger = candidate_tokens if smaller is target_tokens else target_tokens
    return len(smaller & larger) / len(smaller) if smaller else 0.0


def token_similarity_score(target_tokens: set[str], candidate_tokens: set[str]) -> float:
    if not target_tokens or not candidate_tokens:
        return 0.0
    target_core = {token for token in target_tokens if token not in {"da", "de", "do", "das", "dos", "e", "a", "o", "em"}}
    candidate_core = {token for token in candidate_tokens if token not in {"da", "de", "do", "das", "dos", "e", "a", "o", "em"}}
    if not target_core or not candidate_core:
        return 0.0
    intersection = len(target_core & candidate_core)
    return intersection / max(len(target_core), len(candidate_core))


def parse_bool(value: str, default: bool = False) -> bool:
    if value is None:
        return default
    return value.strip().lower() in {"1", "true", "sim", "yes", "y"}


@dataclass
class Record:
    index: int
    action: str
    person_type: str
    full_name: str
    birthdate: str
    cpf: str
    previous_competition: str = ""

    @property
    def normalized_action(self) -> str:
        return normalize_text(self.action)

    @property
    def normalized_person_type(self) -> str:
        return normalize_text(self.person_type)

    @property
    def requires_identity(self) -> bool:
        return self.normalized_action == "inclusao" and self.normalized_person_type == "atleta"

    @property
    def is_commission(self) -> bool:
        return "comissao tecnica" in self.normalized_person_type


@dataclass
class RequestFile:
    source_name: str
    source_url: str
    protocol: str
    team_name: str
    competition_name: str
    pix_receipt_url: str
    records: list[Record]
    raw_text: str


@dataclass
class TeamLinkInfo:
    name: str
    url: str


@dataclass
class TeamSpreadsheetUpdate:
    team_name: str
    total_athletes: int
    pix_receipt_url: str = ""


@dataclass
class RecordResult:
    index: int
    action: str
    person_type: str
    full_name: str
    birthdate: str
    cpf: str
    status: str
    message: str


class DuplicateAthleteError(ValueError):
    pass


class PortabilityMatchError(LookupError):
    pass


class RemovalValidationError(ValueError):
    pass


@dataclass
class PortabilityMatch:
    checkbox: WebElement | None
    label: str
    score: float
    threshold: float
    matched: bool


class AppConfig:
    def __init__(self, path: Path) -> None:
        parser = configparser.ConfigParser()
        if not parser.read(path, encoding="utf-8"):
            raise FileNotFoundError(f"Arquivo de configuracao nao encontrado: {path}")

        self.path = path
        self.parser = parser
        self.username = parser.get("ifut", "username")
        self.password = parser.get("ifut", "password")
        self.login_url = parser.get("ifut", "login_url")
        self.championship_url = parser.get("ifut", "championship_url")
        self.teams_url = parser.get("ifut", "teams_url")
        self.championship_public_id = parser.get("ifut", "championship_public_id", fallback="")
        self.championship_public_slug = parser.get("ifut", "championship_public_slug", fallback="")
        self.team_urls = dict(parser.items("teams")) if parser.has_section("teams") else {}
        self.headless = parse_bool(parser.get("selenium", "headless", fallback="false"))
        self.timeout = parser.getint("selenium", "timeout_seconds", fallback=20)
        self.download_dir = Path(parser.get("drive", "download_dir", fallback=str(DEFAULT_DOWNLOAD_DIR)))
        self.processed_dir = Path(parser.get("drive", "processed_dir", fallback=str(DEFAULT_PROCESSED_DIR)))
        self.failed_dir = Path(parser.get("drive", "failed_dir", fallback=str(DEFAULT_FAILED_DIR)))
        self.results_dir = Path(parser.get("drive", "results_dir", fallback="downloads\\resultados"))
        self.folder_embed_url = parser.get("drive", "folder_embed_url")
        self.service_account_json = Path(parser.get("drive", "service_account_json", fallback="google-service-account.json"))
        self.log_path = Path(parser.get("app", "log_path", fallback=str(DEFAULT_LOG_PATH)))
        self.pause_after_action = parse_bool(parser.get("app", "pause_after_action", fallback="true"), default=True)
        self.dry_run = parse_bool(parser.get("app", "dry_run", fallback="true"), default=True)
        self.wait_between_records_seconds = parser.getfloat("app", "wait_between_records_seconds", fallback=2.0)
        self.team_page_wait_seconds = parser.getfloat("app", "team_page_wait_seconds", fallback=2.0)
        self.portability_popup_wait_seconds = parser.getfloat("app", "portability_popup_wait_seconds", fallback=2.0)
        self.portability_cancel_delay_seconds = parser.getfloat("app", "portability_cancel_delay_seconds", fallback=1.0)
        self.removal_click_delay_seconds = parser.getfloat("app", "removal_click_delay_seconds", fallback=3.0)
        self.removal_confirm_delay_seconds = parser.getfloat("app", "removal_confirm_delay_seconds", fallback=3.0)
        self.portability_source_championship = parser.get("app", "portability_source_championship", fallback="2º COPA AMERICA 2026")
        self.spreadsheet_id = parser.get("sheets", "spreadsheet_id", fallback="")
        self.sheets_range_times = parser.get("sheets", "range_times", fallback="A2:C")
        self.sheets_update_enabled = parse_bool(parser.get("sheets", "update_enabled", fallback="false"), default=False)


class SelectorConfig:
    def __init__(self, path: Path) -> None:
        parser = configparser.ConfigParser()
        if not parser.read(path, encoding="utf-8"):
            raise FileNotFoundError(f"Arquivo de seletores nao encontrado: {path}")
        self.parser = parser

    def get(self, section: str, option: str) -> str:
        value = self.parser.get(section, option, fallback="").strip()
        if not value:
            raise KeyError(f"Seletor ausente: [{section}] {option}")
        return value

    def get_optional(self, section: str, option: str) -> str | None:
        value = self.parser.get(section, option, fallback="").strip()
        return value or None


class DriveTxtDownloader:
    def __init__(self, folder_embed_url: str, download_dir: Path, service_account_json: Path | None = None) -> None:
        self.folder_embed_url = folder_embed_url
        self.download_dir = download_dir
        self.service_account_json = service_account_json
        self.remote_files_by_name: dict[str, str] = {}
        self._drive_service = None
        self._root_folder_id: str | None = None
        self._entry_folder_id: str | None = None
        self._processed_folder_id: str | None = None
        self._failed_folder_id: str | None = None

    def sync(self) -> list[Path]:
        self.download_dir.mkdir(parents=True, exist_ok=True)
        if self.service_account_json and self.service_account_json.exists():
            return self._sync_with_service_account()
        html = self._fetch_html(self.folder_embed_url)
        entry_folder_url = self._find_entry_folder_url(html)
        if entry_folder_url:
            html = self._fetch_html(self._folder_view_url(entry_folder_url))
        matches = re.findall(
            r'<div class="flip-entry"[^>]*>.*?<a href="https://drive\.google\.com/file/d/([^"/]+)/view[^"]*".*?<div class="flip-entry-title">(.*?)</div>',
            html,
            re.DOTALL,
        )
        files: list[Path] = []
        for file_id, title in matches:
            safe_name = title.strip()
            if not safe_name.lower().endswith(".txt"):
                continue
            destination = self.download_dir / safe_name
            content = requests.get(
                f"https://drive.google.com/uc?export=download&id={file_id}",
                timeout=30,
            )
            content.raise_for_status()
            destination.write_bytes(content.content)
            files.append(destination)
        return files

    def _sync_with_service_account(self) -> list[Path]:
        credentials = service_account.Credentials.from_service_account_file(
            str(self.service_account_json),
            scopes=["https://www.googleapis.com/auth/drive"],
        )
        service = build("drive", "v3", credentials=credentials)
        self._drive_service = service
        root_folder_id = self._extract_folder_id(self.folder_embed_url)
        self._root_folder_id = root_folder_id
        entry_folder_id = self._find_named_folder_id(service, root_folder_id, "Entrada")
        self._entry_folder_id = entry_folder_id or root_folder_id
        self._processed_folder_id = self._find_named_folder_id(service, root_folder_id, "Processados")
        self._failed_folder_id = self._find_named_folder_id(service, root_folder_id, "Falhas")
        folder_id = self._entry_folder_id
        query = f"'{folder_id}' in parents and trashed = false and (name contains '.txt' or mimeType = 'text/plain')"
        response = service.files().list(
            q=query,
            fields="files(id, name)",
            pageSize=200,
        ).execute()
        files: list[Path] = []
        for item in response.get("files", []):
            destination = self.download_dir / item["name"].strip()
            self.remote_files_by_name[destination.name] = item["id"]
            request = service.files().get_media(fileId=item["id"])
            buffer = io.BytesIO()
            downloader = MediaIoBaseDownload(buffer, request)
            done = False
            while not done:
                _, done = downloader.next_chunk()
            destination.write_bytes(buffer.getvalue())
            files.append(destination)
        return files

    def _fetch_html(self, url: str) -> str:
        response = requests.get(url, timeout=30)
        response.raise_for_status()
        return response.text

    def _find_entry_folder_url(self, html: str) -> str | None:
        match = re.search(
            r'<a href="(https://drive\.google\.com/drive/folders/[^"]+)".*?<div class="flip-entry-title">Entrada</div>',
            html,
            re.DOTALL,
        )
        return match.group(1) if match else None

    def _folder_view_url(self, folder_url: str) -> str:
        parsed = urlparse(folder_url)
        folder_id = parsed.path.rstrip("/").split("/")[-1]
        return f"https://drive.google.com/embeddedfolderview?id={folder_id}#list"

    def _extract_folder_id(self, url: str) -> str:
        parsed = urlparse(url)
        if "id=" in url:
            match = re.search(r"[?&]id=([^&#]+)", url)
            if match:
                return match.group(1)
        return parsed.path.rstrip("/").split("/")[-1]

    def _find_named_folder_id(self, service, root_folder_id: str, folder_name: str) -> str | None:
        query = (
            f"'{root_folder_id}' in parents and trashed = false "
            f"and mimeType = 'application/vnd.google-apps.folder' and name = '{folder_name}'"
        )
        response = service.files().list(
            q=query,
            fields="files(id, name)",
            pageSize=10,
        ).execute()
        files = response.get("files", [])
        return files[0]["id"] if files else None

    def ensure_status_folders(self, logger=None) -> None:
        """Cria as pastas Processados/Falhas no Drive quando ainda nao existirem."""
        if not self._drive_service or not self._root_folder_id:
            return
        for attribute, folder_name in (("_processed_folder_id", "Processados"), ("_failed_folder_id", "Falhas")):
            if getattr(self, attribute):
                continue
            created = self._drive_service.files().create(
                body={
                    "name": folder_name,
                    "mimeType": "application/vnd.google-apps.folder",
                    "parents": [self._root_folder_id],
                },
                fields="id",
            ).execute()
            setattr(self, attribute, created["id"])
            if logger:
                logger.info("Pasta %s criada no Drive", folder_name)

    def move_remote_file(self, file_name: str, destination_kind: str, logger=None) -> None:
        if not self._drive_service:
            if logger:
                logger.warning("Servico Google Drive nao inicializado; skip movimentacao remota para %s", file_name)
            return
        file_id = self.remote_files_by_name.get(file_name)
        if not file_id:
            if logger:
                logger.warning("Arquivo %s nao encontrado no rastreamento remoto; talvez nao foi baixado via service account", file_name)
            return
        destination_folder_id = self._processed_folder_id if destination_kind == "processed" else self._failed_folder_id
        source_folder_id = self._entry_folder_id
        if not destination_folder_id or not source_folder_id:
            if logger:
                logger.error("Pastas de destino nao configuradas: processados=%s, falhas=%s", self._processed_folder_id, self._failed_folder_id)
            return
        try:
            self._drive_service.files().update(
                fileId=file_id,
                addParents=destination_folder_id,
                removeParents=source_folder_id,
                fields="id, parents",
            ).execute()
            if logger:
                logger.info("Arquivo movido no Drive: %s -> %s", file_name, destination_kind)
        except Exception as exc:
            if logger:
                logger.error("Erro ao mover arquivo no Drive %s: %s", file_name, exc)


class TxtRequestParser:
    HEADER_PATTERN = re.compile(r"^(.*?):\s*(.*)$")
    RECORD_PATTERN = re.compile(
        r"REGISTRO\s+(\d+)\s*\nACAO:\s*(.*?)\s*\nTIPO:\s*(.*?)\s*\nNOME COMPLETO:\s*(.*?)\s*\nDATA DE NASCIMENTO:\s*(.*?)\s*\nCPF:\s*(.*?)"
        r"(?:\s*\nCOMPETICAO ANTERIOR:\s*(.*?))?(?=\n\s*\nREGISTRO\s+\d+\s*\n|\Z)",
        re.DOTALL,
    )

    def parse(self, path: Path) -> RequestFile:
        raw_text = path.read_text(encoding="utf-8", errors="ignore")
        compact = raw_text.replace("\r", "")
        header: dict[str, str] = {}
        for line in compact.splitlines():
            line = line.strip()
            if not line:
                continue
            match = self.HEADER_PATTERN.match(line)
            if match:
                header[normalize_text(match.group(1))] = match.group(2).strip()

        records: list[Record] = []
        for record_match in self.RECORD_PATTERN.finditer(compact):
            records.append(
                Record(
                    index=int(record_match.group(1)),
                    action=record_match.group(2).strip(),
                    person_type=record_match.group(3).strip(),
                    full_name=record_match.group(4).strip(),
                    birthdate=record_match.group(5).strip(),
                    cpf=record_match.group(6).strip(),
                    previous_competition=(record_match.group(7) or "").strip(),
                )
            )

        return RequestFile(
            source_name=path.name,
            source_url="",
            protocol=header.get("protocolo", ""),
            team_name=header.get("equipe", ""),
            competition_name=header.get("competicao", ""),
            pix_receipt_url=header.get("comprovante pix", ""),
            records=records,
            raw_text=compact,
        )


class IfutBot:
    def __init__(self, config: AppConfig, selectors: SelectorConfig, logger: logging.Logger) -> None:
        self.config = config
        self.selectors = selectors
        self.logger = logger
        self.driver = self._build_driver()
        self.wait = WebDriverWait(self.driver, self.config.timeout)

    def _build_driver(self) -> webdriver.Chrome:
        options = ChromeOptions()
        if self.config.headless:
            options.add_argument("--headless=new")
        options.add_argument("--start-maximized")
        options.add_argument("--disable-notifications")
        options.add_argument("--lang=pt-BR")
        service = Service(ChromeDriverManager().install())
        return webdriver.Chrome(service=service, options=options)

    def close(self) -> None:
        self.driver.quit()

    def login(self) -> None:
        self.driver.get(self.config.login_url)
        self._fill(self.selectors.get("login", "username_input"), self.config.username)
        self._fill(self.selectors.get("login", "password_input"), self.config.password)
        self._click(self.selectors.get("login", "submit_button"))
        self.wait.until(ec.url_contains("/campeonatos"))

    def open_teams_page(self) -> None:
        self.driver.get(self.config.teams_url)
        self.wait.until(ec.presence_of_element_located(self._locator(self.selectors.get("teams", "page_ready"))))

    def open_team_page(self, team_name: str) -> None:
        team_url = self._team_url_for(team_name)
        if not team_url:
            raise LookupError(f"URL do time nao mapeada: {team_name}")
        self.driver.get(team_url)
        self.wait.until(ec.url_contains("/time/"))
        time.sleep(self.config.team_page_wait_seconds)

    def _team_url_for(self, team_name: str) -> str | None:
        key = team_search_key(team_name)
        for mapped_key, mapped_url in self.config.team_urls.items():
            if normalize_text(mapped_key) == key:
                return mapped_url
        return None

    def process_request(self, request: RequestFile) -> None:
        self.logger.info("Processando protocolo %s do time %s", request.protocol, request.team_name)
        self.open_team_page(request.team_name)
        self._open_roster_section()
        results: list[RecordResult] = []
        for record in request.records:
            self.logger.info(
                "Registro %s - acao=%s tipo=%s nome=%s",
                record.index,
                record.action,
                record.person_type,
                record.full_name,
            )
            try:
                message = self._execute_record(record)
                self.logger.info("Registro %s processado com sucesso", record.index)
                results.append(self._build_record_result(record, "SUCESSO", message))
            except DuplicateAthleteError as exc:
                self.logger.error("Registro %s falhou: %s", record.index, exc)
                results.append(self._build_record_result(record, "FALHA", str(exc)))
            except Exception as exc:
                self.logger.exception("Registro %s falhou: %s", record.index, exc)
                results.append(self._build_record_result(record, "FALHA", str(exc)))
            time.sleep(self.config.wait_between_records_seconds)
        
        self.logger.info("Escrevendo arquivo de resultado...")
        self._write_result_file(request, results)
        
        self.logger.info("Etapa final: Atualizando planilha de controle (etapa adicional, nao obrigatoria)...")
        self._update_spreadsheet(request)
        
        self.logger.info("Processamento do arquivo finalizado com sucesso")

    def update_all_teams_athlete_counts(self) -> None:
        self.logger.info("[PLANILHA] Iniciando rotina independente de atualizacao de quantidade de atletas")
        for mapped_team_name in self.config.team_urls:
            team_name = mapped_team_name.strip()
            self.logger.info("[PLANILHA] Lendo total de atletas do time '%s' no iFut...", team_name)
            try:
                self.open_team_page(team_name)
                self._open_roster_section()
                total_athletes = self._extract_total_athletes()
                if total_athletes is None:
                    self.logger.warning("[PLANILHA] Nao foi possivel obter total de atletas do time '%s'", team_name)
                    continue
                self._update_spreadsheet_entry(
                    TeamSpreadsheetUpdate(team_name=team_name, total_athletes=total_athletes)
                )
            except Exception as exc:
                self.logger.exception(
                    "[PLANILHA] Falha ao atualizar quantidade do time '%s' na rotina independente: %s",
                    team_name,
                    exc,
                )
        self.logger.info("[PLANILHA] Rotina independente de atualizacao de quantidade de atletas finalizada")

    def _open_roster_section(self) -> None:
        tab_selector = self.selectors.get_optional("team", "roster_tab")
        if tab_selector:
            self._click(tab_selector)
        self.wait.until(ec.presence_of_element_located(self._locator(self.selectors.get("team", "roster_ready"))))

    def _execute_record(self, record: Record) -> str:
        action = record.normalized_action
        if action == "inclusao":
            return self._include_person(record)
        if action == "remocao":
            self._remove_person(record)
            return "Remocao executada"
        if action == "portabilidade":
            self._port_person(record)
            return "Portabilidade executada"
        raise ValueError(f"Acao nao suportada: {record.action}")

    def _include_person(self, record: Record) -> str:
        self._open_include_context(record)
        include_selector = (
            self.selectors.get_optional("actions", "commission_include_button")
            if record.is_commission
            else None
        ) or self.selectors.get("actions", "include_button")
        self._click(include_selector)
        dialog_selector = self.selectors.get_optional("messages", "include_dialog")
        if record.is_commission and dialog_selector:
            self._wait_for_include_dialog_or_fail()
        elif record.is_commission:
            time.sleep(self.config.portability_popup_wait_seconds)
        self._fill_person_form(record)
        if self.config.dry_run:
            self.logger.info("DRY RUN ativo: formulario preenchido, sem clicar em salvar.")
            self._restore_default_team_tab(record)
            return "Formulario preenchido em DRY RUN"
        self._click(self.selectors.get("actions", "save_button"))
        duplicate_message = self._wait_for_include_outcome()
        if duplicate_message:
            self.logger.error("Inclusao recusada para %s: %s", record.full_name, duplicate_message)
            self._close_include_popup()
            self._restore_default_team_tab(record)
            raise DuplicateAthleteError(duplicate_message)
        self.logger.info("Inclusao enviada com sucesso para %s", record.full_name)
        time.sleep(3)
        self._restore_default_team_tab(record)
        return "Inclusao enviada com sucesso"

    def _remove_person(self, record: Record) -> None:
        self._open_removal_context(record)
        row = self._find_person_row(record.full_name)
        time.sleep(self.config.removal_click_delay_seconds)
        try:
            self._click_child(row, self.selectors.get("actions", "remove_row_button"))
        except NoSuchElementException:
            delete_button = row.find_element(
                By.XPATH,
                ".//button[contains(@class, 'text-negative') or .//i[contains(@class, 'delete')]]",
            )
            self.driver.execute_script("arguments[0].click();", delete_button)
        confirm = self.selectors.get_optional("actions", "confirm_remove_button")
        if confirm:
            self.wait.until(ec.presence_of_element_located(self._locator(confirm)))
            time.sleep(self.config.removal_confirm_delay_seconds)
            self._validate_removal_confirmation(record)
            self._click(confirm)
            time.sleep(self.config.removal_confirm_delay_seconds)
        self._restore_default_team_tab(record)

    def _port_person(self, record: Record) -> None:
        self._open_roster_section()
        self._click(self.selectors.get("actions", "portability_button"))
        time.sleep(self.config.portability_popup_wait_seconds)
        self._select_portability_source_championship(record)
        match = self._find_portability_checkbox(record.full_name)
        if not match.matched or match.checkbox is None:
            self._close_portability_popup()
            raise PortabilityMatchError(
                f"Atleta nao encontrado na lista de portabilidade: {record.full_name}. "
                f"Nome mais proximo encontrado: {match.label or 'NENHUM'}. "
                f"Similaridade: {match.score:.2f}. "
                f"Limiar minimo: {match.threshold:.2f}."
            )
        self.driver.execute_script("arguments[0].click();", match.checkbox)
        if self.config.dry_run:
            self.logger.info(
                "DRY RUN ativo: atleta marcado para portabilidade (%s, score=%.2f), sem clicar em inscrever.",
                match.label,
                match.score,
            )
            self._close_portability_popup()
            return
        self._click(self.selectors.get("actions", "portability_submit_button"))
        time.sleep(2)

    def _select_portability_source_championship(self, record: Record) -> None:
        championship_name = record.previous_competition or self.config.portability_source_championship
        self._click(self.selectors.get("actions", "portability_championship_select"))
        option_xpath = self.selectors.get("actions", "portability_championship_option").format(championship_name=championship_name)
        time.sleep(self.config.portability_popup_wait_seconds)
        self._click(option_xpath)
        time.sleep(self.config.portability_popup_wait_seconds)
        self.wait.until(ec.presence_of_element_located(self._locator(self.selectors.get("actions", "portability_player_checkbox"))))

    def _find_portability_checkbox(self, full_name: str) -> PortabilityMatch:
        from difflib import SequenceMatcher

        normalized_target = normalize_text(full_name)
        target_tokens = text_tokens(full_name)
        target_tokens_all = text_tokens_all(full_name)
        threshold = 0.88
        relaxed_threshold = 0.58

        checkboxes = self.driver.find_elements(*self._locator(self.selectors.get("actions", "portability_player_checkbox")))
        best_match = None
        best_label = ""
        best_score = 0.0

        for checkbox in checkboxes:
            label = checkbox.get_attribute("aria-label") or checkbox.text
            normalized_label = normalize_text(label)

            if normalized_target == normalized_label:
                return PortabilityMatch(checkbox, label, 1.0, threshold, True)

            candidate_tokens = text_tokens(label)
            candidate_tokens_all = text_tokens_all(label)
            token_score = token_overlap_score(target_tokens, candidate_tokens)
            token_score_all = token_overlap_score(target_tokens_all, candidate_tokens_all)
            containment_score = token_containment_score(target_tokens_all, candidate_tokens_all)
            similarity_score = token_similarity_score(target_tokens_all, candidate_tokens_all)
            sequence_score = SequenceMatcher(None, normalized_target, normalized_label).ratio()
            combined_score = (token_score * 0.25) + (token_score_all * 0.15) + (containment_score * 0.15) + (similarity_score * 0.20) + (sequence_score * 0.25)

            self.logger.debug(
                "Portabilidade: %s vs %s | token=%.2f seq=%.2f combined=%.2f",
                full_name,
                label,
                token_score,
                sequence_score,
                combined_score,
            )

            if combined_score > best_score:
                best_score = combined_score
                best_match = checkbox
                best_label = label

        if best_score >= threshold or (best_score >= relaxed_threshold and self._looks_like_particle_only_difference(normalized_target, normalize_text(best_label))):
            self.logger.info(
                "Portabilidade: Atleta encontrado com fuzzy matching (score=%.2f): %s",
                best_score,
                full_name,
            )
            return PortabilityMatch(best_match, best_label, best_score, threshold, True)

        if best_label:
            return PortabilityMatch(None, best_label, best_score, threshold, False)

        return PortabilityMatch(None, "", 0.0, threshold, False)

    def _looks_like_particle_only_difference(self, normalized_target: str, normalized_candidate: str) -> bool:
        from difflib import SequenceMatcher

        target_tokens = [token for token in normalized_target.split() if token]
        candidate_tokens = [token for token in normalized_candidate.split() if token]
        particles = {"da", "de", "do", "das", "dos", "e", "a", "o", "em"}
        target_core = [token for token in target_tokens if token not in particles]
        candidate_core = [token for token in candidate_tokens if token not in particles]
        if target_core == candidate_core:
            return True
        if len(target_core) != len(candidate_core):
            return False
        differences = 0
        for target_token, candidate_token in zip(target_core, candidate_core):
            if target_token == candidate_token:
                continue
            if SequenceMatcher(None, target_token, candidate_token).ratio() >= 0.8:
                differences += 1
            else:
                return False
        return differences <= 1 and abs(len(target_tokens) - len(candidate_tokens)) <= 2

    def _validate_removal_confirmation(self, record: Record) -> None:
        try:
            confirmation_msg = self.driver.find_element(
                By.XPATH,
                ".//div[contains(@class, 'q-dialog__message')]"
            )
            dialog_text = confirmation_msg.text.strip()
            
            extracted_name = self._extract_name_from_removal_dialog(dialog_text)
            
            if not extracted_name or not self._names_match(record.full_name, extracted_name):
                self.logger.error(
                    "Nome no popup de confirmacao nao bate: esperado '%s', encontrado '%s'",
                    record.full_name,
                    extracted_name or "NENHUM"
                )
                deny_button = self.driver.find_element(
                    By.XPATH,
                    ".//button[contains(., 'Eita, não')]"
                )
                self.driver.execute_script("arguments[0].click();", deny_button)
                raise RemovalValidationError(
                    f"Nome do registro ({record.full_name}) nao encontrado no popup de confirmacao. "
                    f"Remocao foi cancelada por seguranca."
                )
            
            self.logger.info("Confirmacao de remocao validada: '%s'", extracted_name)
        except NoSuchElementException:
            self.logger.warning("Nao foi possivel validar popup de confirmacao de remocao")

    def _extract_name_from_removal_dialog(self, dialog_text: str) -> str:
        import re
        match = re.search(r"excluir o atleta\s+(.+?)\s+desse", dialog_text, re.IGNORECASE)
        if match:
            return match.group(1).strip()
        return ""

    def _names_match(self, name1: str, name2: str, threshold: float = 0.85) -> bool:
        from difflib import SequenceMatcher
        normalized1 = self._normalize_name(name1)
        normalized2 = self._normalize_name(name2)
        
        if normalized1 == normalized2:
            return True
        
        similarity = SequenceMatcher(None, normalized1, normalized2).ratio()
        return similarity >= threshold

    def _normalize_name(self, name: str) -> str:
        import unicodedata
        name = name.strip().lower()
        name = unicodedata.normalize('NFKD', name)
        name = name.encode('ascii', 'ignore').decode('ascii')
        return ' '.join(name.split())

    def _close_portability_popup(self) -> None:
        cancel_selector = self.selectors.get_optional("actions", "portability_cancel_button")
        if cancel_selector:
            time.sleep(self.config.portability_cancel_delay_seconds)
            self._click(cancel_selector)
            time.sleep(self.config.portability_cancel_delay_seconds)
            checkbox_selector = self.selectors.get_optional("actions", "portability_player_checkbox")
            if checkbox_selector:
                try:
                    self.wait.until(ec.invisibility_of_element_located(self._locator(checkbox_selector)))
                except TimeoutException:
                    self.logger.warning("Pop-up de portabilidade permaneceu aberto apos clicar em Cancelar.")

    def _select_person_type(self, record: Record) -> None:
        type_selector = self.selectors.get_optional("actions", "person_type_select")
        if not type_selector:
            return
        self._click(type_selector)
        option_key = "commission_type_option" if "comissao tecnica" in record.normalized_person_type else "athlete_type_option"
        self._click(self.selectors.get("actions", option_key))

    def _open_include_context(self, record: Record) -> None:
        if not record.is_commission:
            return
        commission_tab_selector = self.selectors.get_optional("team", "commission_tab")
        if commission_tab_selector:
            self._click(commission_tab_selector)
            time.sleep(self.config.portability_popup_wait_seconds)
            team_ready = self.selectors.get_optional("team", "commission_ready") or self.selectors.get("team", "roster_ready")
            self.wait.until(ec.presence_of_element_located(self._locator(team_ready)))

    def _restore_default_team_tab(self, record: Record) -> None:
        if not record.is_commission:
            return
        self._open_roster_section()

    def _open_removal_context(self, record: Record) -> None:
        if not record.is_commission:
            self._open_roster_section()
            return
        commission_tab_selector = self.selectors.get_optional("team", "commission_tab")
        if commission_tab_selector:
            self._click(commission_tab_selector)
            time.sleep(self.config.portability_popup_wait_seconds)
            team_ready = self.selectors.get_optional("team", "commission_ready") or self.selectors.get("team", "roster_ready")
            self.wait.until(ec.presence_of_element_located(self._locator(team_ready)))

    def _fill_person_form(self, record: Record) -> None:
        image_path = Path("perfil-foto-default.png")
        upload_selector = self.selectors.get_optional("form", "image_input")
        if upload_selector and image_path.exists():
            input_element = self.wait.until(ec.presence_of_element_located(self._locator(upload_selector)))
            input_element.send_keys(str(image_path.resolve()))
            time.sleep(1)
        self._fill(self.selectors.get("form", "full_name_input"), record.full_name)
        if record.requires_identity:
            self._fill(self.selectors.get("form", "birthdate_input"), record.birthdate)
            self._fill(self.selectors.get("form", "cpf_input"), record.cpf)
        if record.is_commission:
            role_selector = self.selectors.get_optional("form", "role_input")
            role_option_selector = self.selectors.get_optional("form", "role_default_option")
            if role_selector and role_option_selector:
                self.logger.info("Clicando em Posição")
                self._click(role_selector)
                time.sleep(0.8)
                self.logger.info("Clicando na opção Treinador")
                self._click(role_option_selector)
                self.logger.info("Aguardando 3 segundos para estabilizar o combobox...")
                time.sleep(3.0)
            document_selector = self.selectors.get_optional("form", "commission_document_input")
            if document_selector and record.cpf and normalize_text(record.cpf) != "nao necessario":
                self._fill(document_selector, record.cpf)

    def _find_duplicate_message(self) -> str | None:
        duplicate_selector = self.selectors.get_optional("messages", "duplicate_rg_card")
        if not duplicate_selector:
            return None
        elements = self.driver.find_elements(*self._locator(duplicate_selector))
        for element in elements:
            text = element.text.strip()
            if text:
                return " ".join(text.split())
        return None

    def _wait_for_include_outcome(self) -> str | None:
        end_time = time.time() + self.config.timeout
        while time.time() < end_time:
            duplicate_message = self._find_duplicate_message()
            if duplicate_message:
                return duplicate_message

            dialog_selector = self.selectors.get_optional("messages", "include_dialog")
            if dialog_selector:
                dialogs = self.driver.find_elements(*self._locator(dialog_selector))
                if not dialogs:
                    return None
            time.sleep(0.3)
        return self._find_duplicate_message()

    def _close_include_popup(self) -> None:
        cancel_selector = self.selectors.get_optional("actions", "cancel_button")
        if not cancel_selector:
            return
        try:
            self._click(cancel_selector)
            dialog_selector = self.selectors.get_optional("messages", "include_dialog")
            if dialog_selector:
                self.wait.until(ec.invisibility_of_element_located(self._locator(dialog_selector)))
            else:
                time.sleep(1)
            self._open_roster_section()
        except Exception:
            self.logger.warning("Nao foi possivel fechar o pop-up de inclusao apos duplicidade.")

    def _wait_for_include_dialog_or_fail(self) -> None:
        dialog_selector = self.selectors.get_optional("messages", "include_dialog")
        if not dialog_selector:
            time.sleep(self.config.portability_popup_wait_seconds)
            return

        end_time = time.time() + self.config.timeout
        while time.time() < end_time:
            dialogs = self.driver.find_elements(*self._locator(dialog_selector))
            if dialogs:
                return
            current_url = self.driver.current_url
            if "/atletas/" in current_url:
                raise RuntimeError(
                    "Fluxo de inclusao de comissao saiu do modal e navegou para a tela de detalhe do atleta/comissao."
                )
            time.sleep(0.2)

        raise TimeoutException("Modal de inclusao nao apareceu apos clicar em Adicionar comissao.")

    def _build_record_result(self, record: Record, status: str, message: str) -> RecordResult:
        return RecordResult(
            index=record.index,
            action=record.action,
            person_type=record.person_type,
            full_name=record.full_name,
            birthdate=record.birthdate,
            cpf=record.cpf,
            status=status,
            message=message,
        )

    def _team_link_for_request(self, request: RequestFile) -> str:
        team_key = normalize_text(request.team_name)
        direct_url = self.config.team_urls.get(team_key)
        if direct_url:
            match = re.search(r"/campeonatos/(\d+)/time/(\d+)", direct_url)
            if match:
                team_id = match.group(2)
                if self.config.championship_public_id and self.config.championship_public_slug:
                    return f"https://campeonato.ifut.com.br/c/{self.config.championship_public_id}-{self.config.championship_public_slug}/time/{team_id}"
                else:
                    public_slug = normalize_text(request.competition_name).replace(" ", "-")
                    return f"https://campeonato.ifut.com.br/c/{self.config.championship_public_id}-{public_slug}/time/{team_id}"
            return direct_url
        return self.config.championship_url.rstrip("/")

    def _write_result_file(self, request: RequestFile, results: list[RecordResult]) -> None:
        self.config.results_dir.mkdir(parents=True, exist_ok=True)
        timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
        result_name = f"{Path(request.source_name).stem}-resultado-{timestamp}.txt"
        result_path = self.config.results_dir / result_name
        inclusions = [
            result
            for result in results
            if result.status == "SUCESSO"
            and normalize_text(result.action) in {"inclusao", "portabilidade"}
        ]
        lines = [
            "RESULTADO DA AUTOMACAO AEUV",
            "",
            f"PROTOCOLO: {request.protocol}",
            f"COMPETICAO: {request.competition_name}",
            f"EQUIPE: {request.team_name}",
            f"QUANTIDADE DE ATLETAS INSCRITOS: {len(inclusions)}",
        ]
        if inclusions:
            lines.extend(["ATLETAS INSCRITOS", "-----------------"])
            for result in inclusions:
                lines.append(f"- {result.full_name}")
        lines.extend(["", "RESULTADOS", "----------", ""])
        for result in results:
            lines.extend(
                [
                    f"REGISTRO {result.index:02d}",
                    f"ACAO: {result.action}",
                    f"TIPO: {result.person_type}",
                    f"NOME COMPLETO: {result.full_name}",
                    f"DATA DE NASCIMENTO: {result.birthdate}",
                    f"CPF: {result.cpf}",
                    f"STATUS: {result.status}",
                    f"MENSAGEM: {result.message}",
                    "",
                ]
            )
        lines.append("")
        lines.append("----------------------------------------")
        lines.append("IMPORTANTE:")
        lines.append("Representante da equipe, confira no aplicativo se as informacoes cadastradas estao corretas.")
        lines.append(f"Link do time para conferencia: {self._team_link_for_request(request)}")
        lines.append(f"Gerado em: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}")
        result_path.write_text("\n".join(lines), encoding="utf-8")

    def _extract_total_athletes(self) -> int | None:
        """Extrai o total de atletas da página."""
        try:
           self.logger.debug("[PLANILHA] Procurando elemento com total de atletas...")
           element = self.driver.find_element(By.XPATH, "//div[contains(@class, 'text-h5')]")
           text = element.text
           match = re.search(r'Total de atletas:\s*(\d+)', text)
           if match:
               total = int(match.group(1))
               self.logger.info("[PLANILHA] Total de atletas extraido: %d", total)
               return total
           self.logger.warning("[PLANILHA] Formato de total de atletas nao reconhecido: %s", text)
        except Exception as exc:
           self.logger.warning("[PLANILHA] Nao foi possivel extrair total de atletas: %s", exc)
        return None

    def _find_team_row_in_sheet(self, worksheet, team_name: str) -> tuple[int | None, str | None]:
        normalized_team = normalize_text(team_name)
        search_tokens = [token for token in normalized_team.split() if token]
        best_contains_row = None
        best_contains_value = None

        for idx, row in enumerate(worksheet.iter_rows(values_only=False), 1):
           cell_value = row[0].value
           if not cell_value:
               continue

           normalized_cell = normalize_text(str(cell_value))
           if not normalized_cell:
               continue

           if normalized_cell == normalized_team:
               return idx, str(cell_value)

           if normalized_team in normalized_cell or normalized_cell in normalized_team:
               if best_contains_row is None:
                   best_contains_row = idx
                   best_contains_value = str(cell_value)
               continue

           cell_tokens = [token for token in normalized_cell.split() if token]
           if search_tokens and all(token in cell_tokens for token in search_tokens):
               if best_contains_row is None:
                   best_contains_row = idx
                   best_contains_value = str(cell_value)

        return best_contains_row, best_contains_value

    def _update_spreadsheet_entry(self, entry: TeamSpreadsheetUpdate) -> None:
        """Atualiza uma linha da planilha de controle."""
        if not self.config.sheets_update_enabled or not self.config.spreadsheet_id:
           self.logger.info("[PLANILHA] Atualizacao de planilha desabilitada no config.ini")
           return

        try:
           import tempfile
           from openpyxl import load_workbook

           self.logger.info("[PLANILHA] Conectando ao Google Drive...")
           credentials = service_account.Credentials.from_service_account_file(
               str(self.config.service_account_json),
               scopes=["https://www.googleapis.com/auth/drive"],
           )
           drive_service = build("drive", "v3", credentials=credentials)

           self.logger.info("[PLANILHA] Baixando arquivo: %s", self.config.spreadsheet_id)
           file_id = self.config.spreadsheet_id
           request_obj = drive_service.files().get_media(fileId=file_id)

           with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
               tmp_path = Path(tmp.name)
               downloader = MediaIoBaseDownload(tmp, request_obj)
               done = False
               while not done:
                   _, done = downloader.next_chunk()
           self.logger.info("[PLANILHA] Arquivo baixado: %s", tmp_path)

           self.logger.info("[PLANILHA] Abrindo workbook...")
           wb = load_workbook(tmp_path)
           ws = wb.active

           team_row, matched_team_name = self._find_team_row_in_sheet(ws, entry.team_name)

           if team_row is None:
               self.logger.warning("[PLANILHA] Time '%s' nao encontrado na planilha", entry.team_name)
               tmp_path.unlink(missing_ok=True)
               return

           self.logger.info(
               "[PLANILHA] Time '%s' encontrado na linha %d da planilha como '%s'",
               entry.team_name,
               team_row,
               matched_team_name,
           )
           self.logger.info("[PLANILHA] Atualizando coluna C (QTDA JOGADORES) = %d", entry.total_athletes)
           ws[f"C{team_row}"].value = entry.total_athletes

           if entry.pix_receipt_url:
               self.logger.info("[PLANILHA] Adicionando comprovante PIX na coluna J")
               current_cell = ws[f"J{team_row}"]
               current_value = str(current_cell.value) if current_cell.value else ""
               if current_value and current_value.strip():
                   new_value = f"{current_value}\n{entry.pix_receipt_url}"
                   self.logger.info("[PLANILHA] Comprovante PIX adicionado ao historico existente")
               else:
                   new_value = entry.pix_receipt_url
                   self.logger.info("[PLANILHA] Novo comprovante PIX registrado")
               current_cell.value = new_value
           else:
               self.logger.info("[PLANILHA] Nenhum comprovante PIX para adicionar")

           self.logger.info("[PLANILHA] Salvando workbook em arquivo temporario...")
           wb.save(tmp_path)

           self.logger.info("[PLANILHA] Enviando arquivo de volta para o Drive...")
           drive_service.files().update(
               fileId=file_id,
               media_body=MediaFileUpload(str(tmp_path), mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
           ).execute()

           tmp_path.unlink(missing_ok=True)
           self.logger.info(
               "[PLANILHA] SUCESSO - Planilha atualizada: time='%s' linha=%d atletas=%d",
               entry.team_name,
               team_row,
               entry.total_athletes,
           )

        except FileNotFoundError as exc:
           self.logger.error("[PLANILHA] Erro: Arquivo de servico ou arquivo temporario nao encontrado: %s", exc)
        except Exception as exc:
           self.logger.error("[PLANILHA] Erro ao atualizar planilha (nao afeta o resultado geral): %s", exc)
           self.logger.exception("[PLANILHA] Stack trace:")

    def _update_spreadsheet(self, request: RequestFile) -> None:
        """Atualiza a planilha de controle com dados do processamento.
         
        Esta é uma etapa adicional que NÃO afeta o sucesso geral da automação.
        Erros na atualização são registrados mas não interrompem o fluxo.
        """
        self.logger.info("[PLANILHA] Iniciando atualização da planilha de controle...")
        total_athletes = self._extract_total_athletes()
        if total_athletes is None:
           self.logger.warning("[PLANILHA] Nao foi possivel extrair quantidade de atletas. Abortando atualizar.")
           return
        self._update_spreadsheet_entry(
           TeamSpreadsheetUpdate(
               team_name=request.team_name,
               total_athletes=total_athletes,
               pix_receipt_url=request.pix_receipt_url,
           )
        )

    def _find_person_row(self, full_name: str) -> WebElement:
        rows = self.driver.find_elements(*self._locator(self.selectors.get("team", "roster_row")))
        normalized_name = normalize_text(full_name)
        search_tokens = [token for token in normalized_name.split() if token]
        search_key = " ".join(search_tokens[:2]) if len(search_tokens) >= 2 else normalized_name
        for row in rows:
            row_text = normalize_text(row.text)
            if normalized_name in row_text:
                return row
            if search_key and search_key in row_text:
                return row
            if self._row_matches_name(row_text, normalized_name):
                return row
        raise LookupError(f"Pessoa nao encontrada na lista do time: {full_name}")

    def _row_matches_name(self, row_text: str, normalized_name: str) -> bool:
        from difflib import SequenceMatcher

        row_tokens = text_tokens_all(row_text)
        name_tokens = text_tokens_all(normalized_name)
        if not row_tokens or not name_tokens:
            return False

        row_core = [token for token in row_tokens if token not in {"da", "de", "do", "das", "dos", "e", "a", "o", "em"}]
        name_core = [token for token in name_tokens if token not in {"da", "de", "do", "das", "dos", "e", "a", "o", "em"}]

        if len(row_core) != len(name_core):
            return False

        score = 0
        for expected, candidate in zip(name_core, row_core):
            if expected == candidate:
                score += 1
            elif SequenceMatcher(None, expected, candidate).ratio() >= 0.85:
                score += 1

        return score == len(name_core) and len(name_core) >= 2

    def _fill(self, selector: str, value: str, clear: bool = True) -> None:
        element = self._wait_for_any_selector(selector)
        if clear:
            element.clear()
        element.send_keys(value)

    def _click(self, selector: str) -> None:
        self._guard_against_modal_context_leak(selector)
        element = self.wait.until(ec.element_to_be_clickable(self._locator(selector)))
        self.driver.execute_script("arguments[0].click();", element)

    def _click_child(self, parent: WebElement, selector: str) -> None:
        self._guard_against_modal_context_leak(selector)
        by, value = self._locator(selector)
        child = parent.find_element(by, value)
        self.driver.execute_script("arguments[0].click();", child)

    def _locator(self, selector: str) -> tuple[str, str]:
        if "=" not in selector:
            raise ValueError(f"Seletor invalido: {selector}")
        strategy, value = selector.split("=", 1)
        strategy = strategy.strip().lower()
        value = value.strip()
        mapping = {
            "css": By.CSS_SELECTOR,
            "xpath": By.XPATH,
            "id": By.ID,
            "name": By.NAME,
        }
        if strategy not in mapping:
            raise ValueError(f"Estrategia de seletor nao suportada: {strategy}")
        return mapping[strategy], value

    def _wait_for_any_selector(self, selector: str) -> WebElement:
        selector_parts = [part.strip() for part in selector.split("||") if part.strip()]
        if len(selector_parts) <= 1:
            return self.wait.until(ec.presence_of_element_located(self._locator(selector)))

        end_time = time.time() + self.config.timeout
        last_error: Exception | None = None
        while time.time() < end_time:
            for selector_part in selector_parts:
                try:
                    element = self.driver.find_element(*self._locator(selector_part))
                    if element:
                        return element
                except Exception as exc:
                    last_error = exc
            time.sleep(0.2)

        raise TimeoutException(f"Nenhum seletor encontrado: {selector}") from last_error

    def _guard_against_modal_context_leak(self, selector: str) -> None:
        dialog_selector = self.selectors.get_optional("messages", "include_dialog")
        if not dialog_selector:
            return
        dialogs = self.driver.find_elements(*self._locator(dialog_selector))
        if not dialogs:
            return
        safe_markers = ("cancelar", "salvar", "nome", "posição", "posicao", "documento", "rg", "escolher imagem")
        blocked_markers = ("adicionar", "importar", "remove", "elenco", "comissão", "comissao")
        selector_text = selector.lower()
        if any(marker in selector_text for marker in safe_markers):
            return
        if any(marker in selector_text for marker in blocked_markers):
            raise RuntimeError("Tentativa de clicar fora do modal enquanto o pop-up de inclusao ainda esta aberto.")


def configure_logging(log_path: Path) -> logging.Logger:
    log_path.parent.mkdir(parents=True, exist_ok=True)
    logger = logging.getLogger("ifut_bot")
    logger.setLevel(logging.INFO)
    logger.handlers.clear()
    formatter = logging.Formatter("%(asctime)s [%(levelname)s] %(message)s")

    file_handler = logging.FileHandler(log_path, encoding="utf-8")
    file_handler.setFormatter(formatter)
    logger.addHandler(file_handler)

    console_handler = logging.StreamHandler(sys.stdout)
    console_handler.setFormatter(formatter)
    logger.addHandler(console_handler)
    return logger


def move_file(source: Path, destination_dir: Path) -> None:
    destination_dir.mkdir(parents=True, exist_ok=True)
    target = destination_dir / source.name
    if target.exists():
        target.unlink()
    source.replace(target)


def iter_txt_files(directory: Path) -> Iterable[Path]:
    return sorted(path for path in directory.glob("*.txt") if path.is_file())


def write_default_config() -> None:
    if not DEFAULT_CONFIG_PATH.exists():
        DEFAULT_CONFIG_PATH.write_text(
            """[ifut]
username = seu_email@ifut.com
password = sua_senha
login_url = https://admin.ifut.com.br/login
championship_url = https://admin.ifut.com.br/campeonatos/131038
teams_url = https://admin.ifut.com.br/campeonatos/131038/times
 
[drive]
folder_embed_url = https://drive.google.com/embeddedfolderview?id=10hhnvDF_J7C0LE5JU9BrPST9D8Rf9qk1#list
download_dir = downloads
processed_dir = downloads\\processados
failed_dir = downloads\\falhas
 
[selenium]
headless = false
timeout_seconds = 20
 
[app]
log_path = logs\\ifut.log
pause_after_action = true
dry_run = true
 
[sheets]
spreadsheet_id = 16Tt-7abmpY2CKtY4TQUrqzkFl9T48MFa
update_enabled = false
""",
            encoding="utf-8",
        )

    if not DEFAULT_SELECTORS_PATH.exists():
        DEFAULT_SELECTORS_PATH.write_text(
            """[login]
username_input = css=input[type='email']
password_input = css=input[type='password']
submit_button = css=button[type='submit']

[teams]
page_ready = css=body
search_input =
team_card = xpath=//a[contains(@href, '/campeonatos/131038/time/') or contains(@href, '/time/')]

[team]
roster_tab =
roster_ready = css=body
roster_row = xpath=//*[self::tr or self::div][contains(translate(., 'ABCDEFGHIJKLMNOPQRSTUVWXYZ', 'abcdefghijklmnopqrstuvwxyz'), 'atleta') or contains(translate(., 'ABCDEFGHIJKLMNOPQRSTUVWXYZ', 'abcdefghijklmnopqrstuvwxyz'), 'comissao')]

[actions]
include_button =
portability_button =
person_type_select =
athlete_type_option =
commission_type_option =
search_person_input =
search_person_button =
select_first_search_result =
save_button =
remove_row_button =
confirm_remove_button =

[form]
full_name_input =
birthdate_input =
cpf_input =
role_input =
""",
            encoding="utf-8",
        )


def build_argument_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Automacao iFut para inclusao, remocao e portabilidade")
    parser.add_argument("--config", default=str(DEFAULT_CONFIG_PATH), help="Arquivo INI de configuracao")
    parser.add_argument("--selectors", default=str(DEFAULT_SELECTORS_PATH), help="Arquivo INI de seletores Selenium")
    parser.add_argument("--sync-drive-only", action="store_true", help="Baixa os TXTs do Google Drive e encerra")
    parser.add_argument("--process-local-only", action="store_true", help="Processa somente os TXTs locais ja baixados")
    parser.add_argument("--login-only", action="store_true", help="Abre o iFut e para logo apos o login")
    parser.add_argument(
        "--update-all-team-counts",
        action="store_true",
        help="Atualiza somente a quantidade de atletas de todos os times configurados na planilha",
    )
    parser.add_argument(
        "--analisar-sumulas",
        action="store_true",
        help="Baixa as sumulas digitais do Drive e gera o rascunho das notas oficiais disciplinares",
    )
    parser.add_argument(
        "--ia",
        action="store_true",
        help="Com --analisar-sumulas: gera a nota com IA (Gemini/OpenAI) em vez das regras fixas",
    )
    parser.add_argument(
        "--gerar-pdf-nota",
        metavar="ALVO",
        help="Regera o PDF de uma nota oficial a partir do TXT: caminho do TXT, numero da nota (ex.: 4) "
             "ou protocolo da sumula (SUM-...)",
    )
    parser.add_argument(
        "--gerar-pdf-regulamento",
        metavar="ARQUIVO",
        help="Gera o PDF do regulamento (layout AEUV, com assinatura do Presidente) a partir do arquivo texto, "
             "ex.: regulamento-7-super-liga-união-2026.txt (procurado tambem na pasta regulamento)",
    )
    return parser


def main() -> int:
    write_default_config()
    args = build_argument_parser().parse_args()
    if args.gerar_pdf_regulamento:
        from nota_pdf import ConfigPdf
        from regulamento_pdf import gerar_pdf_regulamento

        config = AppConfig(Path(args.config))
        gerar_pdf_regulamento(args.gerar_pdf_regulamento, ConfigPdf.carregar(Path(args.config)),
                              logger=configure_logging(config.log_path))
        return 0
    if args.gerar_pdf_nota:
        from nota_pdf import regerar_pdf

        config = AppConfig(Path(args.config))
        regerar_pdf(args.gerar_pdf_nota, Path(args.config), configure_logging(config.log_path))
        return 0
    if args.analisar_sumulas:
        from sumula_disciplinar import processar_sumulas

        processar_sumulas(Path(args.config), local_only=args.process_local_only, usar_ia=args.ia)
        return 0
    config = AppConfig(Path(args.config))
    selectors = SelectorConfig(Path(args.selectors))
    logger = configure_logging(config.log_path)

    downloader = DriveTxtDownloader(config.folder_embed_url, config.download_dir, config.service_account_json)
    if not args.process_local_only:
        downloaded = downloader.sync()
        logger.info("Arquivos sincronizados do Drive: %s", len(downloaded))
        if args.sync_drive_only:
            return 0

    bot = IfutBot(config, selectors, logger)
    try:
        bot.login()
        if args.login_only:
            logger.info("Login executado com sucesso; encerrando por --login-only.")
            return 0
        if args.update_all_team_counts:
            bot.update_all_teams_athlete_counts()
            return 0
        parser = TxtRequestParser()
        txt_files = list(iter_txt_files(config.download_dir))
        if not txt_files:
            logger.info("Nenhum arquivo TXT encontrado em %s", config.download_dir)
            return 0
        for txt_file in txt_files:
            try:
                request = parser.parse(txt_file)
                bot.process_request(request)
                move_file(txt_file, config.processed_dir)
                downloader.move_remote_file(txt_file.name, "processed", logger)
            except Exception as exc:
                logger.exception("Falha ao processar %s: %s", txt_file.name, exc)
                move_file(txt_file, config.failed_dir)
                downloader.move_remote_file(txt_file.name, "failed", logger)
    finally:
        bot.close()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

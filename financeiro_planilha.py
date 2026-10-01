"""Leitura dos lancamentos financeiros direto da planilha do Google Drive.

A planilha "AEUV - Financeiro" (aba "Movimentacoes") e alimentada pelo modulo
Financeiro do sistema interno (apps-scripts/sistema-interno). Aqui ela e
exportada como XLSX pela API do Drive e convertida em LancamentoFinanceiro,
sem exigir a API do Sheets habilitada no projeto.
"""

from __future__ import annotations

import configparser
import io
import logging
from datetime import date, datetime
from pathlib import Path

from financeiro_pdf import LancamentoFinanceiro

LOGGER = logging.getLogger(__name__)

NOME_PLANILHA_PADRAO = "AEUV - Financeiro"
NOME_ABA_PADRAO = "Movimentacoes"
MIME_PLANILHA = "application/vnd.google-apps.spreadsheet"
MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _config_financeiro(config_path: Path) -> tuple[str, str, str]:
    parser = configparser.ConfigParser()
    parser.read(config_path, encoding="utf-8")
    parser.read(Path(config_path).with_name("config.local.ini"), encoding="utf-8")
    return (
        parser.get("financeiro", "planilha_id", fallback="").strip(),
        parser.get("financeiro", "planilha_nome", fallback=NOME_PLANILHA_PADRAO).strip(),
        parser.get("financeiro", "aba", fallback=NOME_ABA_PADRAO).strip(),
    )


def _localizar_planilha(service, planilha_id: str, nome: str) -> str:
    if planilha_id:
        return planilha_id
    nome_q = nome.replace("'", "\\'")
    resp = service.files().list(
        q=f"name = '{nome_q}' and mimeType = '{MIME_PLANILHA}' and trashed = false",
        fields="files(id, name, modifiedTime)",
        orderBy="modifiedTime desc",
        pageSize=5,
        supportsAllDrives=True,
        includeItemsFromAllDrives=True,
    ).execute()
    arquivos = resp.get("files", [])
    if not arquivos:
        raise FileNotFoundError(
            f"Planilha '{nome}' nao encontrada no Drive. Abra o modulo Financeiro do sistema interno "
            "ao menos uma vez ou informe [financeiro] planilha_id no config.ini."
        )
    return arquivos[0]["id"]


def _texto_data(valor) -> str:
    if isinstance(valor, datetime):
        return valor.strftime("%d/%m/%Y")
    if isinstance(valor, date):
        return valor.strftime("%d/%m/%Y")
    return str(valor or "").strip()


def _numero(valor) -> float:
    if isinstance(valor, (int, float)):
        return float(valor)
    texto = str(valor or "").replace("R$", "").strip()
    if "," in texto:
        texto = texto.replace(".", "").replace(",", ".")
    try:
        return float(texto)
    except ValueError:
        return 0.0


def carregar_lancamentos_da_planilha(
    config_path: Path, logger: logging.Logger | None = None
) -> list[LancamentoFinanceiro]:
    """Baixa a planilha do Drive e devolve os lancamentos da aba Movimentacoes."""
    import drive_auth
    from googleapiclient.discovery import build
    from googleapiclient.http import MediaIoBaseDownload
    from openpyxl import load_workbook

    log = logger or LOGGER
    planilha_id, nome, aba_nome = _config_financeiro(config_path)
    service = build("drive", "v3", credentials=drive_auth.credenciais_por_arquivo(config_path, log),
                    cache_discovery=False)
    planilha_id = _localizar_planilha(service, planilha_id, nome)
    log.info("[FINANCEIRO] Lendo planilha %s (id=%s, aba=%s)", nome, planilha_id, aba_nome)

    buffer = io.BytesIO()
    requisicao = service.files().export_media(fileId=planilha_id, mimeType=MIME_XLSX)
    baixador = MediaIoBaseDownload(buffer, requisicao)
    concluido = False
    while not concluido:
        _, concluido = baixador.next_chunk()
    buffer.seek(0)

    livro = load_workbook(buffer, read_only=True, data_only=True)
    if aba_nome not in livro.sheetnames:
        raise ValueError(f"Aba '{aba_nome}' nao encontrada na planilha. Abas: {', '.join(livro.sheetnames)}")
    aba = livro[aba_nome]

    lancamentos: list[LancamentoFinanceiro] = []
    for linha in aba.iter_rows(min_row=2, values_only=True):
        celulas = list(linha) + [None] * (14 - len(linha))
        if not str(celulas[0] or "").strip():
            continue
        lancamentos.append(LancamentoFinanceiro(
            id_lancamento=str(celulas[0]).strip(),
            data_movimentacao=_texto_data(celulas[1]),
            tipo=str(celulas[2] or "Entrada").strip(),
            origem=str(celulas[3] or "Geral / Administrativo").strip(),
            categoria=str(celulas[4] or "").strip(),
            descricao=str(celulas[5] or "").strip(),
            valor=abs(_numero(celulas[6])),
            favorecido_pagador=str(celulas[7] or "").strip(),
            documento=str(celulas[8] or "").strip(),
            id_transacao_bancaria=str(celulas[9] or "").strip(),
            emenda=str(celulas[10] or "").strip(),
            comprovante_url=str(celulas[11] or "").strip(),
            responsavel=str(celulas[13] or "").strip(),
        ))
    livro.close()

    log.info("[FINANCEIRO] %s lancamento(s) carregado(s) da planilha", len(lancamentos))
    return lancamentos
